from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)

C_BLUE  = RGBColor(0x1F, 0x49, 0x7D)
C_BLUE2 = RGBColor(0x2E, 0x74, 0xB5)
C_GRAY  = RGBColor(0x50, 0x50, 0x50)
C_RED   = RGBColor(0xC7, 0x25, 0x4E)
C_GREEN = RGBColor(0x10, 0x7C, 0x10)

def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = C_BLUE

def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = C_BLUE2

def h3(text):
    p = doc.add_heading(text, level=3)
    p.runs[0].font.color.rgb = C_GRAY

def body(text, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = bold
    if color:
        r.font.color.rgb = color

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(1)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9)
    r.font.color.rgb = C_RED
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    p._p.get_or_add_pPr().append(shd)

def note(text):
    p = doc.add_paragraph()
    r = p.add_run('※ ' + text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = C_GRAY
    r.italic = True

def sp():
    doc.add_paragraph()

def tbl(headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        r = c.paragraphs[0].runs[0]
        r.bold = True
        r.font.size = Pt(10)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri+1].cells[ci]
            c.text = val
            c.paragraphs[0].runs[0].font.size = Pt(10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)

# ═══════════════════════════════════════════════════
# 표지
# ═══════════════════════════════════════════════════
title = doc.add_heading('MES LLM API', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = C_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('ChromaDB · RAG · 온톨로지 소스 설명')
r.font.size = Pt(16)
r.font.color.rgb = C_BLUE2

sp()
desc = doc.add_paragraph()
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = desc.add_run('자연어 → SQL 변환 시스템의 지식 검색 아키텍처')
r2.font.size = Pt(11)
r2.font.color.rgb = C_GRAY

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 1. 전체 흐름
# ═══════════════════════════════════════════════════
h1('1. 전체 처리 흐름')
body('사용자가 자연어로 질문하면 아래 4단계를 거쳐 SQL이 생성됩니다.')
sp()

tbl(
    ['단계', '모듈', '역할'],
    [
        ('① 개념 매핑',   'ontology_loader.py',  '한국어 개념어 → 관련 테이블/컬럼 식별'),
        ('② 벡터 검색',   'retriever.py (ChromaDB)', '의미적으로 유사한 스키마 청크 top-6 검색'),
        ('③ JOIN 경로',   'fk_resolver.py',      'Dijkstra 알고리즘으로 최단 JOIN 경로 계산'),
        ('④ SQL 생성',    'ollama.py (Qwen2.5)',  '스키마 컨텍스트 기반 SQL 생성'),
    ],
    col_widths=[3, 4.5, 8]
)
sp()

body('처리 파이프라인 (retriever.py → retrieve_context 함수)', bold=True)
sp()
code('사용자 질문: "홀드 상태인 LOT 목록 보여줘"')
code('        ↓')
code('① ConceptEngine   "홀드" → CIM_LOT, CIM_LOT.ISHOLD')
code('        ↓')
code('② ChromaDB RAG    top-6 유사 스키마 청크 반환')
code('        ↓')
code('③ FKPathResolver  식별된 테이블 간 JOIN 경로 계산')
code('        ↓')
code('  [컨텍스트] 스키마청크 + 개념매핑 + JOIN힌트')
code('        ↓')
code('④ Ollama LLM      SELECT * FROM CIM_LOT WITH (NOLOCK) WHERE ISHOLD = \'Y\'')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 2. ChromaDB
# ═══════════════════════════════════════════════════
h1('2. ChromaDB — 벡터 데이터베이스')
body('파일: core/retriever.py')
sp()

h2('2-1. 역할')
body('mes_schema.json의 테이블/컬럼 정보를 벡터로 변환하여 저장합니다. 질문이 들어오면 의미적으로 가장 가까운 스키마 정보를 빠르게 검색합니다.')
sp()

h2('2-2. 시딩 (최초 1회 실행)')
body('서버 시작 시 _seed_schema() 함수가 실행되어 아래 형태의 문서를 생성합니다.')
sp()
body('테이블 단위 문서:', bold=True)
code('[테이블] CIM_LOT')
code('설명: 생산 로트 (PK: LOTID, SITEID)')
code('컬럼: LOTID(varchar(40)), LOTNAME(nvarchar(160)), ISHOLD(varchar(1)), ...')
code('FK: EQUIPMENTID → CIM_EQUIPMENT.EQUIPMENTID')
sp()
body('컬럼 단위 문서:', bold=True)
code('[컬럼] CIM_LOT.ISHOLD (varchar(1))')
code('설명: 홀드여부 (Y/N)')
sp()
note('all-MiniLM-L6-v2 임베딩 모델(약 90MB)로 문서를 벡터화하여 chroma_db/ 폴더에 저장합니다.')
sp()

h2('2-3. 자동 재시딩')
body('mes_schema.json 파일이 변경되면 서버 재시작 시 자동으로 재시딩합니다.')
sp()
code('현재 파일 MD5 해시 계산')
code('        ↓')
code('chroma_db/.schema_hash 에 저장된 이전 해시와 비교')
code('        ↓')
code('다르면 → 컬렉션 삭제 후 재시딩')
code('같으면 → 기존 컬렉션 그대로 사용')
sp()

h2('2-4. 검색')
body('retrieve_context() 함수에서 ChromaDB에 질문을 던져 유사 스키마 청크를 가져옵니다.')
sp()
code('_collection.query(')
code('    query_texts=["홀드 상태인 LOT"],')
code('    n_results=6   # RAG_TOP_K 설정값')
code(')')
code('# → 의미적으로 가장 유사한 테이블/컬럼 문서 6개 반환')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 3. RAG
# ═══════════════════════════════════════════════════
h1('3. RAG — 하이브리드 검색 시스템')
body('파일: core/retriever.py  →  retrieve_context() 함수')
sp()

h2('3-1. RAG란?')
body('RAG(Retrieval-Augmented Generation)는 LLM이 답변을 생성하기 전에 외부 지식을 검색해서 프롬프트에 주입하는 방식입니다. 이 프로젝트에서는 MES DB 스키마 정보를 검색해서 LLM이 정확한 SQL을 생성하도록 돕습니다.')
sp()

h2('3-2. 단순 RAG vs 하이브리드 RAG')
sp()
tbl(
    ['항목', '단순 RAG', '이 프로젝트 (하이브리드)'],
    [
        ('검색 방법', '벡터 유사도만 사용', '벡터 + 개념어 키워드 + FK 그래프'),
        ('JOIN 처리', 'LLM이 스스로 판단', 'Dijkstra로 최단 경로 계산 후 주입'),
        ('도메인 용어', '인식 불가',         '온톨로지로 "불량률", "홀드" 등 인식'),
        ('정확도',     '보통',              '높음'),
    ],
    col_widths=[3.5, 4.5, 6.5]
)
sp()

h2('3-3. 3단계 파이프라인 상세')
sp()
h3('Step 1 — 개념어 기반 테이블 식별 (ConceptEngine)')
body('질문에서 도메인 키워드를 찾아 관련 테이블/컬럼을 미리 확보합니다.')
sp()
code('질문: "이번 달 불량률이 높은 LOT 보여줘"')
code('  "불량" → CIM_LOT, CIM_LOT.LOSSQTY, CIM_LOT.ORIGINALQTY')
code('  "LOT"  → CIM_LOT, CIM_LOT.LOTID')
sp()

h3('Step 2 — ChromaDB 벡터 검색')
body('질문과 의미적으로 유사한 스키마 청크를 top-6개 가져옵니다. Step 1 결과와 합쳐서 중복 없이 테이블 목록을 구성합니다.')
sp()
code('# 개념 테이블 먼저, ChromaDB 결과로 보완')
code('all_tables = concept_tables + chroma_tables  # 중복 제거')
sp()

h3('Step 3 — FK 경로 기반 JOIN 힌트')
body('식별된 테이블들 사이의 최단 JOIN 경로를 계산해 프롬프트에 추가합니다.')
sp()
code('all_tables = ["CIM_LOT", "CIM_EQUIPMENTCLASS"]')
code('  → CIM_LOT → CIM_EQUIPMENT → CIM_EQUIPMENTCLASS  (2홉)')
code('  → JOIN CIM_EQUIPMENT WITH (NOLOCK)')
code('      ON CIM_LOT.EQUIPMENTID = CIM_EQUIPMENT.EQUIPMENTID')
code('     JOIN CIM_EQUIPMENTCLASS WITH (NOLOCK)')
code('      ON CIM_EQUIPMENT.EQUIPMENTCLASSID = CIM_EQUIPMENTCLASS.EQUIPMENTCLASSID')
sp()

h2('3-4. 최종 컨텍스트 구조')
body('LLM 프롬프트에 주입되는 컨텍스트는 아래 3개 섹션으로 구성됩니다.')
sp()
code('[스키마 청크 1] [테이블] CIM_LOT / 설명: ... / 컬럼: ...')
code('---')
code('[스키마 청크 2] [컬럼] CIM_LOT.LOSSQTY (float) / 설명: 불량수량')
code('--- (6개까지)')
code('')
code('## 개념 매핑 (온톨로지)')
code('관련 테이블: CIM_LOT')
code('관련 컬럼: CIM_LOT.LOSSQTY, CIM_LOT.ORIGINALQTY')
code('')
code('## JOIN 경로 힌트 (Dijkstra 최단경로)')
code('JOIN CIM_EQUIPMENT WITH (NOLOCK) ON CIM_LOT.EQUIPMENTID = CIM_EQUIPMENT.EQUIPMENTID')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 4. 온톨로지
# ═══════════════════════════════════════════════════
h1('4. 온톨로지 — 도메인 지식 표현')
sp()

h2('4-1. 관련 파일')
sp()
tbl(
    ['파일', '역할'],
    [
        ('data/mes_ontology.ttl',     'OWL 온톨로지 (Turtle 형식) — 주 사용'),
        ('data/mes_concept_map.json', 'JSON 폴백 — rdflib 없을 때 사용'),
        ('core/ontology_loader.py',   '엔진 로딩 및 검색 로직'),
    ],
    col_widths=[6, 9]
)
sp()

h2('4-2. OWL 온톨로지 구조 (mes_ontology.ttl)')
body('W3C 표준 OWL 형식의 Turtle 파일입니다. MES 도메인 개념을 클래스/속성으로 정의합니다.')
sp()

h3('클래스 정의 예시')
code('@prefix mfg: <http://neurofactory.skku.ac.kr/mes#> .')
code('')
code('mfg:Lot')
code('    a owl:Class ;')
code('    rdfs:label "생산 로트"@ko ;')
code('    rdfs:comment "제조 공정에서 일괄 처리되는 제품 단위"@ko ;')
code('    mfg:mapsToTable "CIM_LOT" ;      ← DB 테이블 매핑')
code('    mfg:primaryKey  "LOTID" ;')
code('    skos:altLabel   "LOT"@ko , "로트"@ko , "작업 로트"@ko .')
sp()

h3('데이터 속성(컬럼) 정의 예시')
code('mfg:holdStatus')
code('    a owl:DatatypeProperty ;')
code('    rdfs:label   "홀드"@ko ;')
code('    skos:altLabel "보류"@ko , "정지"@ko , "잠금"@ko ;')
code('    mfg:mapsToColumn "CIM_LOT.ISHOLD" .   ← DB 컬럼 매핑')
sp()

h3('관계(JOIN) 정의 예시')
code('mfg:lotRunsOnEquipment')
code('    a owl:ObjectProperty ;')
code('    rdfs:domain mfg:Lot ;')
code('    rdfs:range  mfg:Equipment ;')
code('    mfg:joinCondition')
code('        "CIM_LOT.EQUIPMENTID = CIM_EQUIPMENT.EQUIPMENTID" .')
sp()

h2('4-3. 온톨로지 엔진 (ontology_loader.py)')
body('엔진 선택 로직:', bold=True)
sp()
code('rdflib 설치됨  AND  mes_ontology.ttl 존재')
code('        → OntologyEngine   (SPARQL로 OWL 그래프 쿼리)')
code('아니면')
code('        → FallbackConceptMapper (mes_concept_map.json 직접 읽기)')
sp()

h3('OntologyEngine 내부 동작')
bullet('ttl 파싱 → rdflib 트리플 그래프 구성')
bullet('_build_alias_index(): skos:altLabel → 테이블/컬럼 인덱스 빌드')
bullet('_build_join_index(): ObjectProperty의 mfg:joinCondition → JOIN 조건 수집')
bullet('find_tables(query): 질문에서 alias 매칭 → 관련 테이블 반환')
bullet('find_columns(query): 질문에서 alias 매칭 → 관련 컬럼 반환')
bullet('get_join_hint(tables): 테이블 쌍에 대한 온톨로지 JOIN 조건 반환')
sp()

h3('싱글턴 패턴')
body('서버 시작 시 1회만 초기화하여 메모리에 유지합니다. init_chroma() → get_concept_engine() 순서로 호출됩니다.')
sp()
code('# core/retriever.py  init_chroma() 내부')
code('_concept_engine = get_concept_engine()')
code('# → OntologyEngine 또는 FallbackConceptMapper 반환')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 5. FK 경로 계산
# ═══════════════════════════════════════════════════
h1('5. FK 경로 계산 — Dijkstra JOIN 최적화')
body('파일: core/fk_resolver.py')
sp()

h2('5-1. 역할')
body('여러 테이블이 필요한 질문에서 어떤 순서로 JOIN해야 하는지 자동으로 계산합니다. mes_schema.json의 foreign_keys 정보를 그래프로 구성하고 Dijkstra 알고리즘으로 최단 경로를 찾습니다.')
sp()

h2('5-2. 그래프 구조')
code('CIM_LOT ──[EQUIPMENTID]──► CIM_EQUIPMENT')
code('                                  │')
code('                          [EQUIPMENTCLASSID]')
code('                                  │')
code('                                  ▼')
code('                         CIM_EQUIPMENTCLASS')
sp()
note('FK 관계는 양방향으로 구성됩니다. A→B 경로가 없어도 B→A 방향으로 탐색합니다.')
sp()

h2('5-3. 사용 예시')
sp()
code('resolver.shortest_join_path("CIM_LOT", "CIM_EQUIPMENTCLASS")')
code('')
code('결과:')
code('  홉1: CIM_LOT → CIM_EQUIPMENT')
code('       JOIN CIM_EQUIPMENT WITH (NOLOCK)')
code('         ON CIM_LOT.EQUIPMENTID = CIM_EQUIPMENT.EQUIPMENTID')
code('  홉2: CIM_EQUIPMENT → CIM_EQUIPMENTCLASS')
code('       JOIN CIM_EQUIPMENTCLASS WITH (NOLOCK)')
code('         ON CIM_EQUIPMENT.EQUIPMENTCLASSID = CIM_EQUIPMENTCLASS.EQUIPMENTCLASSID')
sp()

h2('5-4. 제약 사항')
bullet('최대 4홉 이내 경로만 사용 (프롬프트 오염 방지)')
bullet('자기 참조 FK 제외 (예: PARENTLOTID → CIM_LOT 자신)')
bullet('경로 없으면 빈 문자열 반환 (LLM이 자체 판단)')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 6. 데이터 파일 정리
# ═══════════════════════════════════════════════════
h1('6. 데이터 파일 정리')
sp()
tbl(
    ['파일', '역할', '수정 시 효과'],
    [
        ('data/mes_schema.json',     '테이블·컬럼 정의, FK 관계\nChromaDB 시딩 원본',    '서버 재시작 시 자동 재시딩'),
        ('data/mes_ontology.ttl',    'OWL 온톨로지\n개념어↔테이블/컬럼 매핑',           '서버 재시작 시 자동 재로드'),
        ('data/mes_concept_map.json','JSON 폴백 개념 맵',                                'rdflib 없을 때만 사용'),
        ('chroma_db/',               '벡터 DB 저장소 (자동 생성)',                        '삭제 후 재시작 시 재시딩'),
        ('chroma_db/.schema_hash',   '스키마 해시 저장\n(변경 감지용)',                   '삭제 시 다음 시작에 재시딩'),
    ],
    col_widths=[5, 6, 4.5]
)
sp()

h1('7. 관련 소스 파일 목록')
sp()
tbl(
    ['파일', '담당 기능'],
    [
        ('core/retriever.py',        '전체 RAG 파이프라인 오케스트레이션\nChromaDB 시딩·검색·재시딩 자동화'),
        ('core/ontology_loader.py',  'OWL/JSON 개념 엔진 로딩\nOntologyEngine / FallbackConceptMapper'),
        ('core/fk_resolver.py',      'FK 그래프 구성 + Dijkstra JOIN 경로 계산'),
        ('core/ollama.py',           'Ollama LLM 호출 (SQL 생성 / 분석)'),
        ('routers/chat.py',          'API 엔드포인트 + LLM 프롬프트 템플릿'),
        ('data/mes_schema.json',     'MES DB 스키마 정의 (26개 테이블)'),
        ('data/mes_ontology.ttl',    'OWL 온톨로지 (Turtle 형식)'),
        ('data/mes_concept_map.json','JSON 폴백 개념 맵'),
    ],
    col_widths=[5.5, 10]
)

doc.save('MES_LLM_API_아키텍처_설명.docx')
print('완료: MES_LLM_API_아키텍처_설명.docx')
