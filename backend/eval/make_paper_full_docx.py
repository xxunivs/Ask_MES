"""
PAPER_FULL_DRAFT.md → Word (표·그림·본문 포함)

Usage (mes-llm-api 루트):
  python eval/build_paper_complete.py   # figure + MD + Word 일괄
  python eval/make_paper_full_docx.py

Output:
  eval/논문초안_NeuroFactory_NL2SQL.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

EVAL = Path(__file__).resolve().parent
ROOT = EVAL.parent
MD = EVAL / "PAPER_FULL_DRAFT.md"
OUT = EVAL / "논문초안_NeuroFactory_NL2SQL.docx"
FIG_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")

C_BLUE = RGBColor(0x1F, 0x49, 0x7D)
C_BLUE2 = RGBColor(0x2E, 0x74, 0xB5)
C_GRAY = RGBColor(0x50, 0x50, 0x50)

SKIP_HEADING = (
    "절 맺음",
    "부록·각주",
    "4.1절과 연결",
    "hwp",
    "체크리스트",
)


def setup_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(3)
    return doc


def add_heading(doc, level: int, text: str):
    text = text.strip()
    if not text:
        return
    p = doc.add_heading(text, level=min(level, 3))
    colors = {1: C_BLUE, 2: C_BLUE2, 3: C_GRAY}
    if p.runs:
        p.runs[0].font.color.rgb = colors.get(level, C_GRAY)


def add_body(doc, text: str, bold: bool = False):
    text = text.strip()
    if not text:
        return
    p = doc.add_paragraph()
    # **bold** segments
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
            r.font.size = Pt(10.5)
        else:
            part = part.replace("`", "")
            if part:
                r = p.add_run(part)
                r.font.size = Pt(10.5)
                r.bold = bold


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    add_inline(p, text)


def add_numbered(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    add_inline(p, text)


def add_inline(paragraph, text: str):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            r.font.size = Pt(10.5)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.size = Pt(10)
        elif part:
            r = paragraph.add_run(part.replace("`", ""))
            r.font.size = Pt(10.5)


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text.replace("**", ""))
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = C_GRAY


def resolve_figure_path(rel: str) -> Path | None:
    rel = rel.strip().replace("\\", "/")
    for base in (EVAL, ROOT):
        p = (base / rel).resolve()
        if p.is_file():
            return p
    return None


def add_figure(doc, rel_path: str):
    path = resolve_figure_path(rel_path)
    if path is None:
        p = doc.add_paragraph()
        p.add_run(f"[그림 파일 없음: {rel_path}]").font.color.rgb = C_GRAY
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(14.5))
    doc.add_paragraph()


def add_table(doc, headers: list[str], rows: list[list[str]]):
    if not headers:
        return
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h.strip()
        if c.paragraphs[0].runs:
            c.paragraphs[0].runs[0].bold = True
            c.paragraphs[0].runs[0].font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < len(headers):
                c = t.rows[ri + 1].cells[ci]
                c.text = val.strip().replace("**", "")
                if c.paragraphs[0].runs:
                    c.paragraphs[0].runs[0].font.size = Pt(10)
    doc.add_paragraph()


def parse_table_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.strip() for c in line.split("|")]


def is_sep_row(line: str) -> bool:
    return bool(re.match(r"^\|[\s\-:|]+\|$", line.strip()))


def should_skip_heading(text: str) -> bool:
    return any(k in text for k in SKIP_HEADING)


def build_from_markdown(doc: Document, md_text: str):
    lines = md_text.splitlines()
    i = 0
    in_abstract = False
    page_break_next_h1 = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("> "):
            i += 1
            continue

        # Title block (first # line)
        if stripped.startswith("# MES ") and i < 5:
            title = stripped.lstrip("# ").strip()
            t = doc.add_heading(title, 0)
            t.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if t.runs:
                t.runs[0].font.color.rgb = C_BLUE
            i += 1
            continue

        if stripped.startswith("**NeuroFactory"):
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = sub.add_run(stripped.replace("**", ""))
            r.font.size = Pt(12)
            r.font.color.rgb = C_BLUE2
            i += 1
            continue

        if stripped.startswith("**저자") or stripped.startswith("**소속") or stripped.startswith("**실험") or stripped.startswith("**프로젝트"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(stripped.replace("**", "")).font.size = Pt(11)
            i += 1
            continue

        if stripped.startswith("**핵심어"):
            add_body(doc, stripped.replace("**", ""), bold=True)
            i += 1
            continue

        if stripped.startswith("|"):
            headers = parse_table_row(stripped)
            i += 1
            if i < len(lines) and is_sep_row(lines[i]):
                i += 1
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep_row(lines[i]):
                    rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, headers, rows)
            continue

        if stripped.startswith("# "):
            text = stripped[2:].strip()
            if should_skip_heading(text):
                i += 1
                while i < len(lines) and not (
                    lines[i].startswith("# ") and not should_skip_heading(lines[i][2:].strip())
                ):
                    i += 1
                continue
            if text.startswith("1. ") or text == "1. 서론":
                doc.add_page_break()
            if text.startswith("4. ") or text.startswith("5. ") or text.startswith("6. "):
                doc.add_page_break()
            add_heading(doc, 1, text)
            i += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:].strip()
            if "초록" in text:
                in_abstract = True
            if should_skip_heading(text):
                i += 1
                while i < len(lines) and not lines[i].startswith("## ") and not lines[i].startswith("# "):
                    i += 1
                continue
            lvl = 2
            if re.match(r"^\d+\.\d+", text):
                lvl = 2
            add_heading(doc, lvl, text)
            i += 1
            continue

        if stripped.startswith("### "):
            text = stripped[4:].strip()
            if should_skip_heading(text):
                i += 1
                continue
            add_heading(doc, 3, text)
            i += 1
            continue

        if stripped.startswith("**[표") or stripped.startswith("**[그림"):
            add_caption(doc, stripped.replace("**", ""))
            i += 1
            continue

        mfig = FIG_RE.match(stripped)
        if mfig:
            add_figure(doc, mfig.group(2))
            i += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped) and not stripped.startswith("1. Yu"):
            add_numbered(doc, re.sub(r"^\d+\.\s*", "", stripped))
            i += 1
            continue

        # paragraph (merge consecutive non-special lines)
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt == "---" or nxt.startswith(("#", "|", "- ", "**[", ">", "![")):
                break
            if re.match(r"^\d+\.\s", nxt) and "Yu," not in nxt:
                break
            para.append(nxt)
            i += 1
        add_body(doc, " ".join(para))


def main():
    if not MD.exists():
        raise FileNotFoundError(f"먼저 통합본 생성: python eval/merge_paper_sections.py → {MD}")

    from datetime import datetime

    md_text = MD.read_text(encoding="utf-8")
    doc = setup_doc()
    build_from_markdown(doc, md_text)
    out = OUT
    try:
        doc.save(str(out))
    except PermissionError:
        out = EVAL / f"논문초안_NeuroFactory_NL2SQL_{datetime.now():%Y%m%d_%H%M%S}.docx"
        doc.save(str(out))
        print(f"기본 경로 잠김(Word 등에서 열림): {OUT}")
    print(f"완료: {out}")


if __name__ == "__main__":
    main()
