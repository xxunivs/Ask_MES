from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)

def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

def body(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = bold

def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F2F2')
    p._p.get_or_add_pPr().append(shd)

def note(text):
    p = doc.add_paragraph()
    r = p.add_run('※ ' + text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    r.italic = True

def sp():
    doc.add_paragraph()

def table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri+1].cells[ci]
            c.text = val
            c.paragraphs[0].runs[0].font.size = Pt(10)

# ── 표지
title = doc.add_heading('MES LLM API', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('설치 및 실행 가이드')
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

sp()
desc = doc.add_paragraph()
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = desc.add_run('자연어 질문 → SQL 자동 생성 → MES DB 조회 시스템')
r2.font.size = Pt(11)
r2.font.color.rgb = RGBColor(0x50, 0x50, 0x50)

doc.add_page_break()

# ── 1. 개요
h1('1. 시스템 구성 개요')
body('본 시스템은 아래 3개 구성요소가 함께 동작합니다.')
sp()
table(
    ['구성요소', '역할', '포트'],
    [
        ('MES LLM API', '자연어 → SQL 변환 서버 (FastAPI)', '8001'),
        ('Ollama',      '로컬 LLM 엔진 (qwen2.5:7b 모델)', '11434'),
        ('MES API',     'MES DB 조회 서버 (별도 운영 중)',  '외부 URL'),
    ]
)
sp()
note('MES API는 이미 운영 중이므로 별도 설치 불필요. URL과 API 키만 확인하면 됩니다.')

doc.add_page_break()

# ── 2. 사전 설치
h1('2. 사전 설치 (최초 1회)')

h2('2-1. Python 3.11 설치')
bullet('https://www.python.org/downloads/ 접속 후 Windows 64-bit Installer 다운로드')
bullet('설치 화면에서 반드시 체크 → "Add Python to PATH" ✔')
bullet('Install Now 클릭하여 설치 완료')
sp()
body('설치 확인 (PowerShell에서 실행):')
code('python --version')
body('→ Python 3.11.x 이상 출력되면 정상')
sp()
note('Python이 이미 설치되어 있다면 버전 확인 후 넘어가세요.')
sp()

h2('2-2. Ollama 설치 및 모델 다운로드')
bullet('https://ollama.com/download 접속 후 Windows 버전 설치')
bullet('설치 완료 후 작업표시줄에 Ollama 아이콘 생성됨')
bullet('PowerShell에서 아래 명령어로 AI 모델 다운로드 (약 4.7GB, 시간 소요):')
sp()
code('ollama pull qwen2.5:7b')
sp()
bullet('다운로드 완료 후 확인:')
code('ollama list')
body('→ qwen2.5:7b 가 목록에 나오면 완료')
sp()
note('Ollama는 Windows 시작 시 자동 실행됩니다. 작업표시줄에서 확인하세요.')

doc.add_page_break()

# ── 3. 프로젝트 설정
h1('3. 프로젝트 설정')

h2('3-1. 프로젝트 폴더 복사')
bullet('팀 공유 폴더에서 mes-llm-api 폴더 전체를 내 PC로 복사합니다.')
bullet('예시 경로: C:\\NeuroFactory\\mes-llm-api')
sp()
note('venv 폴더는 복사하지 않아도 됩니다. 서버 실행 시 자동 생성됩니다.')
sp()

h2('3-2. 환경 설정 파일 (.env) 수정')
body('mes-llm-api 폴더 안의 .env 파일을 메모장으로 열어 아래 항목을 수정합니다.')
sp()
code('# Ollama 설정 (수정 불필요)')
code('OLLAMA_URL=http://localhost:11434')
code('OLLAMA_MODEL=qwen2.5:7b')
code(' ')
code('# MES DB API 접속 정보  ←  담당자에게 최신 정보 확인')
code('MES_API_URL=https://xxxx.trycloudflare.com')
code('MES_API_KEY=<MES_API_KEY>')
sp()
table(
    ['항목', '설명', '수정 여부'],
    [
        ('MES_API_URL', 'MES DB 서버 주소 (Cloudflare 터널 URL)', '담당자 확인 후 수정'),
        ('MES_API_KEY', 'API 인증 키',                            '담당자 확인 후 수정'),
        ('OLLAMA_URL',  'Ollama 서버 주소',                        '수정 불필요'),
    ]
)
sp()
note('MES_API_URL은 서버 재시작 시 변경될 수 있습니다. 401 오류 발생 시 담당자에게 최신 URL을 받으세요.')

doc.add_page_break()

# ── 4. 서버 실행
h1('4. 서버 실행')

h2('4-1. start.bat 실행')
bullet('mes-llm-api 폴더에서 PowerShell 실행')
bullet('아래 명령어 입력:')
sp()
code('.\\start.bat')
sp()
note('Windows 탐색기에서 start.bat 파일을 더블클릭해도 실행됩니다.')
sp()

h2('4-2. 최초 실행 시 자동 처리 항목')
body('처음 실행 시 아래 작업이 자동으로 진행됩니다. (시간 소요)')
bullet('Python 가상환경 생성')
bullet('패키지 자동 설치 (수 분 소요)')
bullet('ChromaDB 임베딩 모델 다운로드 (약 90MB, 최초 1회)')
bullet('MES 스키마 자동 시딩')
sp()
note('두 번째 실행부터는 위 과정이 생략되어 빠르게 시작됩니다.')
sp()

h2('4-3. 정상 실행 확인')
body('아래 메시지가 출력되면 서버가 정상 실행된 것입니다.')
sp()
code('[Startup] ChromaDB 초기화 중...')
code('[ChromaDB] 기존 컬렉션 로드 완료 (XXX 문서)')
code('[Startup] 완료 — MES LLM API 준비됨')
code('INFO:     Uvicorn running on http://0.0.0.0:8001')

doc.add_page_break()

# ── 5. 테스트
h1('5. 테스트')

h2('5-1. 서버 상태 확인')
body('브라우저에서 아래 주소 접속:')
sp()
table(
    ['용도', 'URL'],
    [
        ('API 테스트 화면 (Swagger UI)', 'http://localhost:8001/docs'),
        ('서버 상태 확인',               'http://localhost:8001/health'),
        ('Ollama 연결 확인',             'http://localhost:8001/chat/health'),
    ]
)
sp()

h2('5-2. Swagger UI로 테스트')
bullet('브라우저에서 http://localhost:8001/docs 접속')
bullet('POST /chat/nl-to-sql 클릭 → "Try it out" 클릭')
bullet('아래 내용 입력 후 Execute 클릭:')
sp()
code('{"question": "LOT 목록 100건 보여줘", "limit": 100}')
sp()
bullet('Response Body에서 sql 과 data 항목 확인')
sp()

h2('5-3. 테스트 질문 예시')
body('아래 질문들로 테스트하세요. (오류 가능성이 낮은 단순 조회)')
sp()
questions = [
    'LOT 목록 100건 보여줘',
    '홀드 상태인 LOT 목록 보여줘',
    '재작업 상태인 LOT 보여줘',
    'LOT별 불량수량이 0보다 큰 것만 보여줘',
    '설비 전체 목록 보여줘',
    '설비 유형별 설비 수를 보여줘',
    '작업지시 전체 목록과 상태를 보여줘',
    '완료 상태가 아닌 작업지시 목록 보여줘',
    '자재 전체 목록 보여줘',
    '안전재고가 있는 자재 목록 보여줘',
    '공정 세그먼트 전체 목록과 소요시간을 보여줘',
    '불량 유형 목록 보여줘',
]
for q in questions:
    bullet(q)

doc.add_page_break()

# ── 6. 오류 해결
h1('6. 자주 발생하는 오류')
sp()
table(
    ['오류', '원인', '해결 방법'],
    [
        ('Ollama 연결 실패',           'Ollama가 실행되지 않음',          '작업표시줄 Ollama 아이콘 확인\n없으면 Ollama 앱 재실행'),
        ('401 Unauthorized',          'API URL 만료 또는 키 오류',        '.env의 MES_API_URL / MES_API_KEY\n담당자에게 최신 정보 확인'),
        ('500 Internal Server Error', 'SQL에 없는 컬럼 사용',             '질문을 더 단순하게 바꿔 재시도'),
        ('Ollama 타임아웃',            '모델 처리 과부하',                 '잠시 후 재시도'),
        ('start.bat 실행 안됨',        'PowerShell 경로 문제',            'PowerShell에서 .\\start.bat 입력'),
    ]
)

sp()

# ── 7. 주요 파일
h1('7. 주요 파일 안내')
sp()
table(
    ['파일 / 폴더', '설명'],
    [
        ('start.bat',            '서버 실행 스크립트'),
        ('.env',                 '환경 설정 (MES_API_URL, MES_API_KEY 수정)'),
        ('data/mes_schema.json', 'MES 테이블/컬럼 정의. 수정 후 재시작하면 자동 반영'),
        ('routers/chat.py',      'API 엔드포인트 및 LLM 프롬프트'),
        ('core/ollama.py',       'Ollama 호출 설정 (타임아웃 등)'),
        ('core/db_client.py',    'MES API 호출 설정'),
        ('chroma_db/',           '벡터 DB 저장 폴더 (자동 생성)'),
        ('venv/',                'Python 가상환경 (공유 불필요, 자동 생성)'),
    ]
)
sp()
note('chroma_db 폴더를 삭제하고 재시작하면 스키마가 처음부터 다시 시딩됩니다.')

doc.save('MES_LLM_API_설치가이드.docx')
print('완료: MES_LLM_API_설치가이드.docx')
